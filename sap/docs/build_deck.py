"""System design deck in the Wilmar Consultancy design system (1920x1080, Fraunces + Plus Jakarta Sans).

Builds HTML, then Chrome prints the PDF; --pptx also writes a PPTX with one picture per slide (from the PDF pages).
Diagrams are made with the arc-draw MCP (draw.io, SAP BTP solution diagram rules) and live in img/diagrams/.

  python3 build_deck.py            # HTML
  ./build_deck.sh                  # HTML, PDF, PPTX
"""
import html, os, sys

S = []
def cover(t, sub, foot): S.append(dict(k="cover", t=t, sub=sub, foot=foot))
def need(eyebrow, t, points): S.append(dict(k="need", e=eyebrow, t=t, points=points))
def table(eyebrow, t, head, rows, widths, pills=None, lead=None): S.append(dict(k="table", e=eyebrow, t=t, head=head, rows=rows, widths=widths, pills=pills, lead=lead))
def diagram(eyebrow, t, img, lead=None, maxw=1664, maxh=540): S.append(dict(k="diagram", e=eyebrow, t=t, img=img, lead=lead, maxw=maxw, maxh=maxh))
def shot(eyebrow, t, lead, img, cap=None): S.append(dict(k="shot", e=eyebrow, t=t, lead=lead, img=img, cap=cap))
def cards(eyebrow, t, items, lead=None): S.append(dict(k="cards", e=eyebrow, t=t, items=items, lead=lead))
def status(eyebrow, t, rows): S.append(dict(k="status", e=eyebrow, t=t, rows=rows))
def close(eyebrow, t, lead, decision, points): S.append(dict(k="close", e=eyebrow, t=t, lead=lead, decision=decision, points=points))

DECK = "Analytics and planning on Fiori"

# ------------------------------------------------------------------ content (plain, calm, three points a slide)
cover("Analytics and planning on Fiori", "System design, features, test results, and how it compares with SAP Analytics Cloud",
      "Prepared for Wilmar  ·  7 October 2026  ·  Working on sample data and on a real S/4HANA Cloud system")

need("Why we built it", "One system for plan and data", [
 ("The need", "Planning and analysis on S/4HANA data today means a second product with its own license, connections and security."),
 ("The question", "How much of the SAC working model can run inside the ERP landscape, under the ERP's own authorizations?"),
 ("The answer so far", "Stories, models, planning with locking and rules, data and multi actions, a calendar, CDS sources and sharing, tested end to end on a real system with two users.")])

table("Current release · 1 of 2", "Build, analyse and see", ["Area", "Capability"], [
 ["Stories", "Pages on a grid; eleven chart types, KPI, tables, input controls, text, images, web pages, feeds, comments, buttons"],
 ["Data Analyser", "Free analysis with rows, columns and filters; calculated measures; CSV and Excel export"],
 ["Models", "Measures with exception aggregation, hierarchies, Date hierarchy, currencies, CSV and CDS sources"],
 ["Insight", "Variance explainer, value driver tree, Compass simulation (Monte Carlo)"],
 ["Navigation", "A floating button opens a menu in the order of the work; usable on a phone"],
 ["Story layout", "Resize from every edge, arrow keys, Tidy up, saved views of tables"]],
 [360, 1304])

table("Current release · 2 of 2", "Plan, automate and share", ["Area", "Capability"], [
 ["Planning", "Editable table, spreading, formulas, copy and paste, versions, publish, history, comments"],
 ["Control", "Data Locking by region, validation rules (error or warning), audit log"],
 ["Data actions", "Seven step types: copy, allocation, formula, currency conversion, scale, delete, embedded"],
 ["Multi actions", "Eleven step types: data action, publish, versions, locking, import, forecast, API, PaPM, comments, copy to model, CDS import"],
 ["Calendar", "Processes and tasks, dependencies, approvals, reminders, Gantt, month, week, day and list views"],
 ["Sharing", "Owner, view, edit and everyone, enforced by the server for stories, models, actions and events"]],
 [360, 1304])

diagram("System design", "The big picture", "diagrams/architecture.png", maxh=610)

cards("System design", "Three design principles", [
 ("Modular", "The library zsac.lib serves any Fiori app. Widgets and data sources are plug-ins registered in one place."),
 ("Metadata is data", "Models, stories, actions and files are rows, edited by the same screens that show them."),
 ("One provider contract", "Mock for development and tests, OData V4 for the backend, CDS views for outside data.")])

diagram("Data", "Live and import sources", "diagrams/sources.png",
        lead="The service aggregates with $apply; if it cannot, the app reads the rows and aggregates itself.")

table("Backend", "ABAP RAP at a glance", ["Part", "Design"], [
 ["Tables", "ZSAC_* tables for files, stories, models, versions, facts, actions, runs, shares and calendar events"],
 ["Business objects", "Managed, strict(2); one request writes a root and its children"],
 ["Actions", "Version: create private, publish, revert. Facts: bulk write and delete"],
 ["Server rules", "Locked versions refuse writes; model structure is checked on save; only the owner shares"],
 ["Access control", "Owner, shared with me and everyone on every projection; the server refuses what the client hides"],
 ["Services", "ZUI_SAC_O4 for the app, ZUI_SAC_BIZ with business views (G/L, sales) as live sources"]],
 [360, 1304])

diagram("Delivery", "Landscape and delivery", "diagrams/landscape.png",
        lead="The same code runs on mock data, in Business Application Studio and deployed to the system.")

status("Operations", "Security and operations", [
 ("ok", "In place", "Calls run under the user's own session; access controls apply to every read, also on the projections."),
 ("ok", "In place", "Owner, view, edit and everyone for stories, models, actions and events; checked with two users."),
 ("ok", "In place", "Locked versions, locking regions and validation rules refuse writes on the server and in every engine."),
 ("ok", "In place", "Every action run is recorded: status, counts, duration, user."),
 ("caution", "Open", "No row level security on plan data; actions run in the browser, not as a server job."),
 ("caution", "Open", "Versions and cell comments are not owned; content moves only through abapGit and the seed class.")])

shot("Navigation", "The floating menu", "No side bar. A button opens the menu above or below it, wherever there is room; it can be dragged out of the way. Pages are grouped in the order of the work: prepare data, plan, report.", "nav-menu-open.jpg")

shot("Real system · G/L", "Live G/L analysis", "A story on the G/L CDS view with 1,092 rows, read live from the S/4HANA Cloud system.", "gl-01-story-overview.jpg", "Story G/L Analysis, page 1, on system my402225")
shot("Real system · G/L", "The model on a CDS view", "The live model maps fields of the service to dimensions and measures. The source panel shows the service, entity and mapping.", "01-modeller-live-cds-source.jpg", "Modeller: live CDS source on the sales view")
shot("Real system · planning", "Versions and edits", "A planner works in a private copy of the budget. Nothing is public until it is published.", "06-planning-versions-dialog.jpg", "Planning: versions dialog")
shot("Real system · planning", "Validation rules", "A limit for plan values. An error refuses the value; a warning lets it through and says so.", "08-planning-validation-rule.jpg", "Rule: max 100 per cell, error")
shot("Real system · planning", "Data locking", "A region locks part of the data. The locked cells are marked in the table.", "10-planning-lock-region-dialog.jpg", "Data Locking dialog")
shot("Real system · actions", "A lock stops the run", "The multi action stops at the failing publish step and says why; the earlier step stays done.", "13-multi-action-lock-stops-run.jpg", "Multi action: data locking stops the publish")

shot("Sample data", "Stories and charts", "Charts, KPIs and input controls on one page, built in the designer.", "03-story-overview.png")
shot("Sample data", "Data actions", "The step flow on the left, the step editor on the right. Validate, trace and run from the toolbar.", "11-data-action-designer.png")
shot("Sample data", "Why did it change?", "A headline, the named drivers and ranked dimensions, with bars to drill into.", "05-variance-explainer.png")

cards("Quality", "How it was tested", [
 ("221 unit tests", "Engines, providers, sharing, locking, calendar, charts and story layout run in Node in about a second. All pass."),
 ("28 end-to-end checks", "Every page on the real system with two accounts, plus phone layout, exports and interface checks. Screenshots for each."),
 ("14 defects fixed", "Found by the checks, not by users: data loss on delete, missing access rules, export noise, dark-on-dark charts, navigation.")])

table("Real system", "What was tested", ["Feature", "Result"], [
 ["Live CDS models, stories with filters, all widget types", "Works"],
 ["Planning: private version, edit, publish, revert, comments", "Works"],
 ["Data locking, validation rules, audit", "Works"],
 ["Data and multi actions; 329 facts; several currencies", "Works"],
 ["Sharing and security with two users", "Works, after fixes"],
 ["CSV and Excel export, input control filter, phone layout", "Works, after fixes"],
 ["Import of 1,092 values in 11 seconds", "Works"],
 ["PaPM step (the system has no PaPM)", "As designed"]],
 [1180, 484], pills={1: ["ok", "ok", "ok", "ok", "ok", "ok", "ok", "neutral"]})

table("Compared with SAC · 1 of 2", "Capability against SAC", ["Capability", "This app", "Real SAC"], [
 ["Stories, analysis, models", "Wide set, eleven charts", "Complete and mature"],
 ["Planning with versions", "Yes, with locking and rules", "Plus broader workflows"],
 ["Data actions", "Seven step kinds, currency conversion", "Plus cross-model, richer allocation"],
 ["Multi actions", "Eleven step types", "Similar, more connectors"],
 ["Value driver tree, Compass", "Yes", "Yes, with more options"],
 ["Calendar and approvals", "Yes, with Gantt and reminders", "Yes, with email notices"],
 ["Sharing and roles", "Owner, view, edit, everyone", "Teams, roles, row level rules"]],
 [620, 520, 524])

table("Compared with SAC · 2 of 2", "Where SAC goes further", ["Capability", "This app", "Real SAC"], [
 ["Predictive and insights", "Four forecast methods, variance", "Smart Predict, smart insights, Just Ask"],
 ["Scripting, custom widgets", "Widget registry (code)", "Analytics Designer scripting"],
 ["Mobile and Office", "Phone layout; CSV and Excel files", "Mobile app, offline, Excel add-in"],
 ["Output", "CSV and Excel", "PDF, PowerPoint, publications, schedules"],
 ["Data connections", "CDS over OData V4", "Many, including BW and Datasphere"],
 ["Scale", "Browser calculations", "In-memory engine and server jobs"],
 ["Transport and governance", "abapGit and seed class", "Content Network, transport, audit tools"]],
 [620, 520, 524])

cards("Compared with SAC", "Where this app is stronger", [
 ("Inside the landscape", "No extra product, license or tenant. It runs in S/4HANA and the Fiori launchpad, on the user's own authorizations."),
 ("Plan data stays in the ERP", "No replication. Same transport and backup path, close to the source."),
 ("Open to change", "Every step, widget and rule can be adapted in days. CDS sources, variance analysis and the calendar tied to planning were built for this use.")])

cards("Compared with SAC", "Where SAC is stronger", [
 ("Wider scope", "Smart Predict, Just Ask, a mobile app, the Excel add-in, scripting and PDF or PowerPoint output are not here."),
 ("Scale and speed", "Calculations run in the browser. Eight steps on 1,092 rows took 85 seconds; large models will be slower."),
 ("Support and governance", "We own the code and its upgrades. Plan data has no row level security, and content transport is basic.")])

table("Gaps", "What it would take to close the gaps", ["Gap to SAC", "Effort", "Approach"], [
 ["Server side data action jobs", "Large", "Run the engine as an ABAP job; keep the browser run for small models"],
 ["Row level security on plan data", "Medium", "Access control on facts by dimension member"],
 ["PDF and PowerPoint output, publications", "Medium", "Server rendering of stories; scheduled mail"],
 ["More data connections", "Medium each", "A source adapter per system behind the provider contract"],
 ["Excel add-in", "Large", "Office add-in on the same OData service"],
 ["Predictive and natural language", "Large", "Connect an AI service; the widget slots exist"],
 ["Content transport", "Medium", "Export and import of stories and models as files"]],
 [620, 220, 824], pills={1: ["caution", "caution", "caution", "caution", "caution", "caution", "caution"]})

cards("Decision", "When to use which", [
 ("Use this app", "Data in S/4HANA CDS views, planning with locking and rules, no second license, and a fast tailored pilot."),
 ("Use real SAC", "Enterprise planning with many users and large volumes, many source systems, or PDF, Excel and mobile needs."),
 ("Use real SAC", "Predictive, natural language and scripting needs, with SAP support and a product roadmap.")])

close("Next steps", "Pilot on the G/L model",
      "The feature set is complete on sample data and verified end to end on a real system with two users.",
      "Decision for Wilmar: run a pilot with one finance team on the G/L model, then decide against SAC from what the users miss.",
      [("1", "Pick the pilot team and the stories and models they share."),
       ("2", "Measure with real volumes and decide if actions need a server job."),
       ("3", "Choose which gaps matter first: row level security, output, connections.")])

# ------------------------------------------------------------------ HTML
e = lambda s: html.escape(s).replace("\n", "<br>")
FONTS = "https://fonts.googleapis.com/css2?family=Fraunces:ital,wght@0,300..700;1,300..700&family=Plus+Jakarta+Sans:wght@400..700&display=swap"
CSS = """
@page { size: 1920px 1080px; margin: 0; }
:root { --s100:#f6f4ee; --s200:#edebe3; --card:#fffdf9; --hair:#dcd9ce; --ink:#0f2a33; --body:#46504f; --muted:#6b6a5e; --deep:#004a5a; --teal:#5294a3; --taupe:#918f80; --coral:#e07a4f; --coraltext:#ad4d26; --ondeep:#f6f4ee; --dark:#073540; --dark2:#0a3d49; --darkcard:#0e4552; }
* { box-sizing: border-box; } html, body { margin: 0; padding: 0; }
body { font-family: 'Plus Jakarta Sans', 'Helvetica Neue', Arial, sans-serif; color: var(--ink); -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.s { position: relative; width: 1920px; height: 1080px; overflow: hidden; page-break-after: always; break-after: page; background: var(--s100); }
.s.dark { background: var(--dark); color: var(--ondeep); }
.wrap { position: absolute; left: 128px; right: 128px; top: 96px; }
.eyebrow { font-size: 24px; letter-spacing: 4px; text-transform: uppercase; color: var(--coraltext); font-weight: 600; }
.dark .eyebrow { color: #f2a27e; }
h1 { font-family: 'Fraunces', Georgia, serif; font-weight: 400; font-size: 72px; line-height: 1.05; margin: 20px 0 0; letter-spacing: -0.5px; color: var(--ink); }
.dark h1 { color: var(--ondeep); }
.lead { font-size: 28px; line-height: 1.4; color: var(--body); margin: 18px 0 0; }
.dark .lead { color: #c8d6d6; }
.footer { position: absolute; left: 128px; right: 128px; bottom: 64px; height: 40px; display: flex; align-items: center; font-size: 24px; color: var(--muted); }
.footer .logo { font-family: 'Fraunces', Georgia, serif; font-size: 28px; color: var(--deep); font-weight: 600; letter-spacing: .5px; display: flex; align-items: center; gap: 12px; }
.footer .logo i { display: inline-block; width: 14px; height: 14px; border-radius: 50%; background: var(--coral); }
.footer .deck { margin-left: 40px; } .footer .pg { margin-left: auto; font-variant-numeric: tabular-nums; }
.dark .footer, .dark .footer .logo { color: #9fb7bb; } .dark .footer .logo { color: var(--ondeep); }
.card { background: var(--card); border: 1px solid var(--hair); border-top: 8px solid var(--deep); border-radius: 28px; padding: 40px 40px 36px; }
.dark .card { background: var(--darkcard); border-color: #1d5a68; border-top-color: var(--teal); }
.card h3 { font-family: 'Fraunces', Georgia, serif; font-weight: 500; font-size: 38px; line-height: 1.15; margin: 0 0 18px; }
.card p { font-size: 28px; line-height: 1.45; margin: 0; color: var(--body); }
.dark .card p { color: #c8d6d6; }
.frame { background: var(--card); border: 1px solid var(--hair); border-radius: 28px; box-shadow: 0 30px 80px rgba(7,53,64,.16); overflow: hidden; }
.frame img { display: block; }
.glow { position: absolute; border-radius: 50%; background: radial-gradient(closest-side, rgba(82,148,163,.30), rgba(82,148,163,0)); }
table { border-collapse: separate; border-spacing: 0; background: var(--card); border: 1px solid var(--hair); border-radius: 28px; overflow: hidden; font-size: 24px; line-height: 1.3; }
th { text-align: left; font-weight: 600; letter-spacing: 3px; text-transform: uppercase; font-size: 24px; color: var(--ondeep); background: var(--deep); padding: 18px 28px; }
td { padding: 11px 28px; border-top: 1px solid var(--hair); color: var(--body); vertical-align: middle; } td:first-child { color: var(--ink); font-weight: 600; }
.pill { display: inline-block; font-size: 24px; font-weight: 600; padding: 6px 22px; border-radius: 999px; white-space: nowrap; }
.pill.ok { background: #d6e8ec; color: #0f4f5e; } .pill.caution { background: #f5e2c4; color: #7a4e0e; } .pill.critical { background: #f6d5c6; color: #8f3517; } .pill.neutral { background: #e3e0d3; color: #4e4d42; }
.row { display: flex; align-items: center; gap: 36px; background: var(--card); border: 1px solid var(--hair); border-radius: 28px; padding: 18px 36px; margin-bottom: 14px; }
.row .pill { width: 170px; text-align: center; flex: none; } .row span.t { font-size: 28px; color: var(--ink); line-height: 1.3; }
.num { font-family: 'Fraunces', Georgia, serif; font-size: 72px; color: var(--coraltext); line-height: 1; }
.callout { background: var(--darkcard); border-left: 10px solid var(--coral); border-radius: 28px; padding: 40px 48px; font-family: 'Fraunces', Georgia, serif; font-size: 44px; line-height: 1.25; color: var(--ondeep); }
.cap { font-size: 24px; color: var(--muted); }
"""

def footer(n):
    return f"<div class='footer'><div class='logo'><i></i>Wilmar Consultancy</div><div class='deck'>{e(DECK)}</div><div class='pg'>{n}</div></div>"

def head(s, width=None):
    w = f"style='width:{width}px'" if width else ""
    lead = f"<div class='lead'>{e(s['lead'])}</div>" if s.get("lead") else ""
    return f"<div class='eyebrow'>{e(s['e'])}</div><h1 {w}>{e(s['t'])}</h1>{lead}"

def render_html():
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{e(DECK)}</title><link rel='preconnect' href='https://fonts.googleapis.com'><link rel='stylesheet' href='{FONTS}'><style>{CSS}</style></head><body>"]
    for n, s in enumerate(S, 1):
        k = s["k"]
        if k == "cover":
            out.append(f"<div class='s dark'><div class='glow' style='left:1000px;top:-200px;width:1300px;height:1300px'></div>"
                       f"<div class='wrap' style='top:300px'><div class='eyebrow'>System design</div><h1 style='font-size:112px;line-height:1.02;margin-top:32px;width:1500px'>{e(s['t'])}</h1><div class='lead' style='font-size:38px;margin-top:36px;width:1300px'>{e(s['sub'])}</div></div>"
                       f"<div class='footer'><div class='logo'><i></i>Wilmar Consultancy</div><div class='deck'>{e(s['foot'])}</div></div></div>")
        elif k == "need":
            cols = "".join(f"<div class='card' style='flex:1'><div class='num'>{i}</div><h3 style='margin-top:20px'>{e(t)}</h3><p>{e(b)}</p></div>" for i, (t, b) in enumerate(s["points"], 1))
            out.append(f"<div class='s'><div class='wrap'>{head(s)}</div><div style='position:absolute;left:128px;right:128px;top:360px;display:flex;gap:40px'>{cols}</div>{footer(n)}</div>")
        elif k == "cards":
            cols = "".join(f"<div class='card' style='flex:1'><h3>{e(t)}</h3><p>{e(b)}</p></div>" for t, b in s["items"])
            top = 400 if s.get("lead") else 350
            out.append(f"<div class='s'><div class='wrap'>{head(s)}</div><div style='position:absolute;left:128px;right:128px;top:{top}px;display:flex;gap:40px'>{cols}</div>{footer(n)}</div>")
        elif k == "table":
            ws = s["widths"]; pills = s.get("pills") or {}
            hd = "<tr>" + "".join(f"<th style='width:{w}px'>{e(h)}</th>" for h, w in zip(s["head"], ws)) + "</tr>"
            rows = ""
            for ri, r in enumerate(s["rows"]):
                cells = ""
                for ci, c in enumerate(r):
                    if ci in pills: cells += f"<td><span class='pill {pills[ci][ri]}'>{e(c)}</span></td>"
                    else: cells += f"<td>{e(c)}</td>"
                rows += f"<tr>{cells}</tr>"
            out.append(f"<div class='s'><div class='wrap'>{head(s)}</div><div style='position:absolute;left:128px;top:300px'><table style='width:{sum(ws)}px'>{hd}{rows}</table></div>{footer(n)}</div>")
        elif k == "diagram":
            from PIL import Image
            iw, ih = Image.open(os.path.join("img", s["img"])).size
            sc = min(s["maxw"] / iw, s["maxh"] / ih); w, h = int(iw * sc), int(ih * sc)
            top = 340 if s.get("lead") else 300
            out.append(f"<div class='s'><div class='wrap'>{head(s)}</div>"
                       f"<div class='frame' style='position:absolute;left:{(1920 - w) // 2}px;top:{top}px;padding:12px'><img src='img/{s['img']}' style='width:{w}px;height:{h}px'></div>"
                       f"{footer(n)}</div>")
        elif k == "shot":
            from PIL import Image
            iw, ih = Image.open(os.path.join("img", s["img"])).size
            w = 992; h = int(w * ih / iw); top = 540 - h // 2 - 30
            cap = f"<div class='cap' style='position:absolute;left:800px;top:{top + h + 28}px;width:992px'>{e(s['cap'])}</div>" if s.get("cap") else ""
            out.append(f"<div class='s'><div class='glow' style='left:620px;top:{top - 180}px;width:1350px;height:{h + 360}px'></div>"
                       f"<div class='wrap' style='width:600px;right:auto;top:300px'><div class='eyebrow'>{e(s['e'])}</div><h1 style='font-size:64px'>{e(s['t'])}</h1><div class='lead' style='font-size:28px'>{e(s['lead'])}</div></div>"
                       f"<div class='frame' style='position:absolute;left:800px;top:{top}px;width:{w}px;height:{h}px'><img src='img/{s['img']}' style='width:{w}px;height:{h}px'></div>{cap}{footer(n)}</div>")
        elif k == "status":
            rows = "".join(f"<div class='row'><span class='pill {p}'>{e(l)}</span><span class='t'>{e(t)}</span></div>" for p, l, t in s["rows"])
            out.append(f"<div class='s'><div class='wrap'>{head(s)}</div><div style='position:absolute;left:128px;right:128px;top:300px'>{rows}</div>{footer(n)}</div>")
        elif k == "close":
            pts = "".join(f"<div style='display:flex;gap:28px;align-items:baseline;margin-bottom:22px'><div class='num' style='font-size:56px;color:#f2a27e;width:48px'>{a}</div><div style='font-size:28px;line-height:1.35;color:#c8d6d6'>{e(b)}</div></div>" for a, b in s["points"])
            out.append(f"<div class='s dark'><div class='glow' style='left:1100px;top:-300px;width:1300px;height:1300px'></div><div class='wrap'>{head(s)}</div>"
                       f"<div style='position:absolute;left:128px;top:400px;width:1000px'>{pts}</div>"
                       f"<div class='callout' style='position:absolute;left:1180px;top:400px;width:612px;font-size:38px'>{e(s['decision'])}</div>{footer(n)}</div>")
    out.append("</body></html>")
    return "".join(out)

def render_pptx(path, pdf_images):
    """One picture per slide (the PDF pages), so the PPTX looks like the PDF."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    for img in pdf_images:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
    prs.save(path)

if __name__ == "__main__":
    open("SAC_Fiori_System_Design.html", "w", encoding="utf-8").write(render_html())
    if "--pptx" in sys.argv:
        import glob
        render_pptx("SAC_Fiori_System_Design.pptx", sorted(glob.glob("/tmp/deckpages/p-*.png")))
    print(len(S), "slides")
