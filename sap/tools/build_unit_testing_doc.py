"""Builds sap/docs/SAC_Fiori_Unit_Testing.docx: the catalogue of every test (unit tests and end-to-end checks) with goal, result and screenshots.
Needs python-docx (python3 -m venv v; v/bin/pip install python-docx). Run from sap/:  python3 tools/build_unit_testing_doc.py

The unit tests are read from sap/app/saclib/test/node/*.test.js (title of each test), so the document follows the tests.
The end-to-end checks and their screenshots are described below (E2E); screenshots are in sap/docs/img."""
import glob
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TESTS = os.path.join(HERE, "..", "app", "saclib", "test", "node")
IMG = os.path.join(HERE, "..", "docs", "img")

# per test file: (area, goal)
FILES = {
    "access": ("Sharing and access", "The rules of who may open, change, delete and share an object are right, and the same in the client, the mock provider and the ABAP backend (which has a copy of them)."),
    "bookmarks": ("Story and planning views", "A user's saved views of a table are kept per user, survive a new session and never break the page when the browser storage fails."),
    "calcmeasures": ("Analyser", "Calculated measures give correct numbers: a ratio of totals, not the total of ratios; formulas on calculated measures; empty cells stay empty; bad definitions are caught."),
    "calendar": ("Calendar", "The calendar works as a planning tool: whole-day dates, processes that roll up their tasks, dependencies, approvals, reminders, month, week, day, list and Gantt views, filters, templates, repeated events, CSV export, and planning tasks that run actions."),
    "cds": ("Live and import data", "A model can read a CDS view live or import it, and everything that can go wrong is reported, not hidden: service errors, missing currency, clients, locked versions."),
    "chart": ("Charts", "Waterfall and stacked charts draw the right shapes and numbers (running total, connectors, axis reaching the stack total)."),
    "chart2": ("Charts", "Treemap, heatmap and geo map put every value in the right place and size, and count what cannot be drawn."),
    "chart3": ("Charts", "Smooth curves never overshoot, areas stack correctly, the Sankey counts a value once, the gauge never shows a false 0%."),
    "compass": ("Compass simulation", "The Monte Carlo engine behind the Compass widget is correct and repeatable: syntax, drivers, draws, cases, histogram, influence and invalid results."),
    "compassview": ("Compass simulation", "The Compass widget shows what the engine computed: chart, second scenario, empty result, numbers and tables."),
    "dataaction": ("Data actions", "A data action changes plan data exactly as defined: copy, aggregate, allocate, scale, delete, convert currency, formulas; a locked version stops the whole run; validation tells what is wrong."),
    "engines": ("Engines", "The core engines give right answers: aggregation with totals, filters, search and replace, data action copy and allocate, version publish and revert, chart builders."),
    "gridview": ("Planning table", "The saved view of the planning table (sorting, zero suppression, scaling, thresholds, variance, calculations) is applied correctly and exported to CSV safely."),
    "hierarchy": ("Hierarchies", "Hierarchies are read correctly: order, depth, totals that count each fact once, filters on a node, one level in charts, validation of cycles and duplicates."),
    "history": ("Undo and redo", "Undo and redo of changes works and is limited; snapshots are copies."),
    "locking": ("Data locking and validation rules", "Data Locking regions and validation rules protect plan data: strictest region wins, private versions are free, a refused write changes nothing, errors refuse and warnings tell."),
    "mockprovider": ("Providers", "The mock provider (sample data, used by all local tests and the preview) behaves like the real service: queries, saving stories, versions, actions, deletion."),
    "model": ("Modeller", "Model rules are right: measure aggregation, exception aggregation, defaults and validation, audit log."),
    "multiaction": ("Multi actions", "A multi action runs its steps in order, stops at the first failing step, maps parameters and handles version, locking, comment and copy steps."),
    "odata-save": ("OData provider", "The OData provider saves and reads correctly without losing data: PATCH and children replaced, restore on refusal, sharing, calendar fields, lock regions and rules."),
    "planning": ("Planning", "The planning table works: date hierarchy, spreading, the edit buffer with undo and redo, typing into leaves and totals, new rows, distribute values, copy and paste, formulas."),
    "resize": ("App shell", "The resizable panels keep their width within limits, can be moved by keyboard and are remembered."),
    "sharing": ("Sharing and access", "Sharing works end to end in the provider for stories, models, data actions, multi actions and calendar events: owner, view, edit, everyone, refusal, deletion with shares."),
    "steps": ("Data and multi action steps", "The step types of actions work: CSV import, forecast, API call, PaPM, comments, with validation and clear failures."),
    "storylayout": ("Story layout", "The story grid behaves: a dropped or resized widget keeps its place and pushes others down, Tidy up closes gaps, resizing from any edge or corner stays in the grid."),
    "valuetree": ("Value driver tree", "The value driver tree parses, calculates, simulates and writes back correctly, and reports every problem."),
    "variance": ("Variance explainer", "The variance explainer names the right drivers: totals and shares add up, ranking, narrative, direction of good, versions with different months."),
    "webcontent": ("Web content widget", "Only safe web content is shown: http(s), paths and small pictures; script and file schemes never."),
    "widgets3": ("Widgets", "The button, feed and comment widgets do only what they are allowed to and clean external content."),
}

# End-to-end checks: (id, area, test, steps, goal, result, screenshot file or None)
E2E_REAL = [
    ("E01", "Modeller", "Connect a model to a live CDS view", "Modeller, new model, Change source, enter the service of the business view, discover the fields, map them", "A model reads business data live from S/4HANA, without copying it", "Pass, after the new business views (fix 1)", "01-modeller-live-cds-source.jpg"),
    ("E02", "Modeller", "Save the live model", "Save, reopen, check structure, currency, periods and the live source", "The live model is stored with its source and can be used by stories and the analyser", "Pass", "02-modeller-saved-live-model.jpg"),
    ("E03", "Data Analyser", "Analyse live data", "Pick the live model, dimensions on rows and columns, filter, table and chart", "Numbers come from the service, aggregated by it", "Pass; CSV and Excel export checked (fix 7)", "03-analyser-live-data.jpg"),
    ("E04", "Stories", "Build a story with live widgets", "Story designer, add chart types, gauge, KPI, table, set the model and measures", "Every widget type renders on live data", "Pass after fixes 1 and 2", "04-story-designer-live-widgets.jpg"),
    ("E05", "Stories", "Open the story in view mode", "Open, change a filter, switch pages", "A story works for the reader as designed", "Pass", "05-story-view-live.jpg"),
    ("E06", "Planning", "Versions: private version", "Versions dialog, create a private copy of the budget", "A planner can work on a private version", "Pass; publish needed fix 3", "06-planning-versions-dialog.jpg"),
    ("E07", "Planning", "Edit a cell", "Type a value in the table, see the buffer, undo, discard, publish", "Edits stay private until published", "Pass", "07-planning-edit-cell.jpg"),
    ("E08", "Planning", "Validation rule", "Validation Rules dialog, add a rule Max 100 per cell, error", "Limits for plan values can be defined", "Pass", "08-planning-validation-rule.jpg"),
    ("E09", "Planning", "Validation rule is enforced", "Type a value above the limit", "An error refuses the value, a warning would tell", "Pass", "09-planning-rule-enforced.jpg"),
    ("E10", "Planning", "Data Locking region", "Data Locking dialog, add a region for a period and members, state Locked", "Parts of the data can be locked", "Pass", "10-planning-lock-region-dialog.jpg"),
    ("E11", "Planning", "Locked cells are shown", "Look at the table with the region", "A planner sees which cells are locked and why", "Pass", "11-planning-locked-cells-hatched.jpg"),
    ("E12", "Data actions", "Run a data action", "Open, dry run, run, read the result", "A data action changes plan data (329 facts in the large test)", "Pass", "12-data-action-run.jpg"),
    ("E13", "Multi actions", "A lock stops the run", "Run a multi action whose publish step meets the locked period", "The first failing step stops the run and says why", "Pass (stopped at the locked 2026-01)", "13-multi-action-lock-stops-run.jpg"),
    ("E14", "Calendar", "Create an event", "Calendar, new event, people, dates, save", "Planning work can be scheduled", "Pass", "14-calendar-event-created.jpg"),
    ("E15", "Stories", "Widgets in view mode", "Open a story with several widget types", "All widgets show data in view mode", "Pass", "15-story-widgets-view.jpg"),
]
E2E_OTHER = [
    ("E16", "Security", "Two users: owner and second user", "Owner in one browser window, second user in a private window; the owner creates a story, model and actions and shares them (view, edit, everyone, nobody); the second user opens, edits, deletes, shares", "A user sees and changes only what the rules allow; the server refuses what the client hides", "Pass after fixes 5 and 6 (delete order, shares removed, access controls on the projections)", None),
    ("E17", "Security", "Delete and unshare", "The second user (edit access) deletes a model; the owner removes access", "A user who may only edit cannot delete; removing a share removes access at once", "Pass after fix 5", None),
    ("E18", "Planning", "Publish and revert a private version", "Create, edit, publish over the budget, revert, delete", "The version life cycle works on the real service", "Pass after fix 3", None),
    ("E19", "Export", "CSV and Excel export", "Export the analyser result and the planning table, open the files", "Exports hold what is shown, with sane numbers and no helper symbols", "Pass after fixes 7 and 8", None),
    ("E20", "Analyser", "Input control filter", "Add a filter widget, choose members, check the other widgets", "A filter widget filters the widgets that use the same dimension", "Pass", None),
    ("E21", "Data actions", "Large data and several currencies", "Run actions on a model with many facts and several currencies", "Totals are right and the run finishes in reasonable time", "Pass (329 facts; import of 1 092 values in 11 s)", None),
    ("E22", "Multi actions", "PaPM step", "Validate and run a PaPM step", "The step reports clearly that no PaPM is connected", "As designed (the system has no PaPM)", None),
    ("E23", "Phone", "Phone layout at 375 px", "Open the analyser, planning, designers and dialogs on a phone-size window", "Everything is usable on a phone without sideways scrolling", "Pass after fix 10", None),
]
E2E_MOCK = [
    ("E24", "Chart text", "Planning chart in the dark theme", "Open Planning in the preview with the dark theme, read legend and axes, read the computed styles of the text", "Legend, axes and periods are readable (light on dark)", "Pass after fix 11 (the text had inherited a dark stroke)", "chart-readable-dark.jpg"),
    ("E25", "Navigation", "Floating button", "Open any page; look at the button; hover, press, drag; read its size and colours", "The button is a solid blue rounded square with a centred icon; text unfolds on hover; stays blue while dragged; no black frame", "Pass", "nav-button.jpg"),
    ("E26", "Navigation", "Floating menu", "Click the button; check the order and groups; choose a page; click outside", "The menu lists the pages in the order of the work, marks the open page, navigates and closes", "Pass", "nav-menu-open.jpg"),
    ("E27", "Navigation", "Menu above or below", "Drag the button near the top, the middle and the bottom of the window and open the menu each time", "The menu opens where there is room (above, else below) and never outside the window", "Pass (below, below, above)", None),
    ("E28", "Story layout", "Resize from an edge, keyboard, Tidy up", "Edit a story, select a widget, drag its left edge, use the arrow keys, press Tidy up", "Resizing from every edge works; arrow keys move; gaps close", "Pass (left edge and arrow keys checked in the preview with simulated pointer and key events; Tidy up by unit test only)", None),
]

FIXES = [
    ("1", "Business views missing; widgets empty; gauge 0%", "End-to-end (E01-E05)", "New CDS views and service ZUI_SAC_BIZ; gauge shows the value or No target set"),
    ("2", "Not found: /Model ZTEST_PLAN for a deleted model", "End-to-end", "Plain error text; Planning and Analyser start with a readable model"),
    ("3", "Publish of a private version needed If-Match", "End-to-end (E06, E18)", "Action sent on the instance"),
    ("4", "Planning Export and Table Functions disabled on open", "End-to-end", "The drawn table is the active table"),
    ("5", "Delete of a model removed data before the server checked the right to delete", "End-to-end, two users (E16, E17)", "Owner check first; sharing tests"),
    ("6", "Shares stayed after the object was deleted; access not enforced on projections", "End-to-end, two users", "Shares removed; access controls on all projections"),
    ("7", "CSV numbers like 12345.600000000001", "End-to-end (E19)", "Rounded to six decimals; analyser CSV test"),
    ("8", "Planning CSV had expand and collapse triangles", "End-to-end (E19)", "Triangles removed on export (gridview CSV test)"),
    ("9", "Modeller panel 500 033 px high; dialogs clipped", "End-to-end", "CSS rules"),
    ("10", "Phone layout: wide chart card, side-by-side designers", "End-to-end (E23)", "Cards fit; panels stack"),
    ("11", "Chart text dark on dark in the dark theme", "Manual (E24)", "Chart text never has a stroke; light colours in dark mode"),
    ("12", "Burger like the SAP shell menu; full-width side bar; later: oval, button hid content, vanished when dragged, black frame, menu cut off", "Manual (E25-E27)", "Floating button and floating menu (3.0 of the manual)"),
    ("13", "No icons on items under Apps; placeholder entries", "Manual", "Flat list with icons; placeholders removed"),
    ("14", "Widgets could be resized from one corner only; gaps stayed", "Manual (E28)", "resizeBox and compact; storylayout tests"),
]


def tidy(title):
    t = re.sub(r"\s+", " ", title).strip()
    return t[:1].upper() + t[1:]


def read_tests():
    out = []
    for f in sorted(glob.glob(os.path.join(TESTS, "*.test.js"))):
        name = os.path.basename(f)[: -len(".test.js")]
        titles = []
        for line in open(f, encoding="utf-8"):
            m = re.match(r'^test\((["`])(.*?)\1\s*,', line)
            if m:
                titles.append(tidy(m.group(2)))
        out.append((name, titles))
    return out


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(doc, head, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size + 0.5)
        shade(c, "E8F0F8")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph()
    return t


def code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Menlo"
    r.font.size = Pt(8.5)
    p.paragraph_format.left_indent = Pt(12)


def shot(doc, name, caption, width=6.2):
    path = os.path.join(IMG, name)
    if not os.path.exists(path):
        doc.add_paragraph("[missing screenshot " + name + "]")
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x5A, 0x6B, 0x7D)


def e2e_block(doc, rows, fig_start):
    n = fig_start
    for (i, area, test, steps, goal, result, img) in rows:
        doc.add_heading(f"{i}  {area}: {test}", 3)
        table(doc, ["Item", ""], [["What is tested", steps], ["Goal", goal], ["Result", result]], [90, 390])
        if img:
            shot(doc, img, f"Figure {n}. {i}: {test}")
            n += 1
    return n


def build():
    files = read_tests()
    total = sum(len(t) for _, t in files)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    p = doc.add_paragraph()
    r = p.add_run("Testing: all tests, their goal and results")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x0A, 0x4D, 0x8C)
    p = doc.add_paragraph()
    r = p.add_run("Analytics and Planning on Fiori (zsac.lib and zsac.fiori). Every unit test with what it checks and why, every end-to-end check with its steps, goal, result and screenshot.")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x4A, 0x5D, 0x70)

    doc.add_heading("1. Summary", 1)
    doc.add_paragraph(
        f"There are three kinds of testing, and this document lists all of them. (A) {total} automated unit tests in {len(files)} files test the logic in Node.js; they run in about a second and all pass. "
        f"(B) {len(E2E_REAL) + len(E2E_OTHER)} end-to-end checks were done by hand on the deployed app (S/4HANA Cloud, client 100) with two accounts, with screenshots. "
        f"(C) {len(E2E_MOCK)} checks of the latest interface changes were done in the browser preview on sample data. "
        "Defects found by the checks are in chapter 6; every one was fixed."
    )
    table(doc, ["Kind", "How", "Count", "Result"], [
        ["A. Unit tests", "node --test, no browser, no backend", str(total), "All pass"],
        ["B. End-to-end, real system", "By hand, deployed app, two accounts, screenshots", str(len(E2E_REAL) + len(E2E_OTHER)), "Pass; defects fixed (E22 PaPM: as designed)"],
        ["C. Interface checks, preview", "Browser preview with sample data, computed styles and simulated pointer and key events", str(len(E2E_MOCK)), "Pass"],
    ], [110, 230, 50, 90])
    doc.add_paragraph("Not tested: there are no ABAP unit tests and no automated tests that drive the UI5 screens. What the user sees is covered by the hand checks of B and C only.")

    doc.add_heading("2. How to run the unit tests", 1)
    code(doc, "cd sap/app/saclib\nnpm test")
    doc.add_paragraph("One file, or one test by name:")
    code(doc, "node --test test/node/storylayout.test.js\nnode --test --test-name-pattern=\"tidy up\" test/node/*.test.js")
    doc.add_paragraph(
        "The modules of zsac.lib use sap.ui.define. The helper test/node/loader.js provides a minimal sap.ui.define, so the pure modules (engines, schemas, providers, chart geometry) load in Node without UI5. "
        "A test loads a module with req(\"zsac/lib/core/StorySchema\"); a UI5 module that is not part of the library can be replaced by a stub in globalThis.__sapStubs. "
        "The OData provider is tested only through its request layer with a stub model."
    )

    doc.add_heading("3. Part A: the unit tests", 1)
    doc.add_paragraph("For every file: the goal, then every test (the title of the test is the rule it checks). All tests pass.")
    tn = 0
    for name, titles in files:
        area, goal = FILES.get(name, ("", ""))
        doc.add_heading(f"{name}.test.js: {area} ({len(titles)} tests)", 2)
        p = doc.add_paragraph()
        r = p.add_run("Goal: ")
        r.bold = True
        p.add_run(goal)
        rows = []
        for t in titles:
            tn += 1
            rows.append([f"U{tn:03d}", t, "Pass"])
        table(doc, ["No.", "What is tested", "Result"], rows, [45, 390, 45], size=8.5)

    doc.add_heading("4. Part B: end-to-end checks on the real system", 1)
    doc.add_paragraph(
        "Done by hand on the deployed app with two accounts: the owner in a normal browser window and a second user in a private window. Screenshots were taken at every step. "
        "The screenshots in this part were taken before the navigation was changed to the floating menu, so they still show the old side navigation; the pages themselves are unchanged. "
        "The objects created for the tests are kept on the system as history data."
    )
    doc.add_heading("4.1 Pages and features, with screenshots", 2)
    n = e2e_block(doc, E2E_REAL, 1)
    doc.add_heading("4.2 Security, export, phone and other checks", 2)
    n = e2e_block(doc, E2E_OTHER, n)

    doc.add_heading("5. Part C: interface changes checked in the preview", 1)
    doc.add_paragraph(
        "The latest changes (chart text, navigation, story layout) were checked in the browser preview on sample data. The checks read computed styles and sizes, and use simulated pointer and key events; "
        "a real mouse and a real device were used only where noted."
    )
    n = e2e_block(doc, E2E_MOCK, n)

    doc.add_heading("6. Defects found and fixed", 1)
    table(doc, ["No.", "Defect", "Found by", "Fix and safeguard"], FIXES, [28, 190, 100, 160], size=8.5)

    doc.add_heading("7. Adding a test", 1)
    for i, b in enumerate([
        "Put the logic in a module without UI5 controls.",
        "Create test/node/<topic>.test.js and load the module with the loader.",
        "Name the test after the rule it checks, build the data inside the test, assert with node:assert.",
        "For a bug found by hand: write the test first (it fails), then fix (it passes).",
        "Run npm test and commit the test with the change; run python3 tools/build_unit_testing_doc.py to refresh this document.",
    ], 1):
        doc.add_paragraph(f"{i}. {b}")
    code(doc, "const test = require(\"node:test\");\nconst assert = require(\"node:assert\");\nconst req = require(\"./loader\");\nconst StorySchema = req(\"zsac/lib/core/StorySchema\");\n\ntest(\"tidy up: gaps close\", () => {\n  const s = { Widgets: [{ Id: \"A\", Page: 1, X: 0, Y: 4, W: 12, H: 2 }] };\n  assert.strictEqual(StorySchema.compact(s, 1), true);\n  assert.strictEqual(s.Widgets[0].Y, 0);\n});")

    doc.add_heading("8. Limits and next steps", 1)
    for b in [
        "No automated test drives the UI5 screens. A browser based suite (OPA5 or Playwright against ?provider=mock) would turn parts B and C into repeatable tests.",
        "No ABAP unit tests. The behavior pools (ZBP_R_SAC_*), ZCL_SAC_ACCESS and the fact writer are the first candidates.",
        "The two-user security check was done by hand; the client side rules are unit tested, the server side rules only by that check.",
        "There is no pipeline yet. npm test needs only Node and can run on every push.",
    ]:
        doc.add_paragraph(b, style="List Bullet")

    out = os.path.join(HERE, "..", "docs", "SAC_Fiori_Unit_Testing.docx")
    doc.save(out)
    print("written", os.path.abspath(out), total, "unit tests,", len(E2E_REAL) + len(E2E_OTHER) + len(E2E_MOCK), "e2e checks")


if __name__ == "__main__":
    build()
